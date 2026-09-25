"""The operations waiting for actual fuel come as a sheet to fill in and upload back.

Actual fuel is recorded weekly or monthly, often into a spreadsheet. The app
hands out the waiting operations already listed - code, vehicle, when, route,
allocation - with the litres column empty. Uploaded back through Impor Massal,
the rows filled in are recorded and the rows still empty are skipped as not
yet filled, instead of being reported as errors.
"""

from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from fuel_predictor.delivery.security import ROUTE_CAPABILITIES
from fuel_predictor.domain.identity import Capability
from fuel_predictor.main import create_app
from tests.test_actual_fuel_outcomes import _operation_with_prediction, _train_baseline


def test_the_waiting_sheet_lists_each_operation_with_an_empty_litres_column(
    tmp_path: Path,
) -> None:
    with TestClient(create_app(database_path=tmp_path / "operations.sqlite3")) as client:
        _train_baseline(client)
        first = _operation_with_prediction(client, 24)["operation"]
        second = _operation_with_prediction(client, 36)["operation"]

        response = client.get("/bahan-bakar-aktual/menunggu.xlsx")

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    sheet = load_workbook(BytesIO(response.content))["Bahan Bakar Aktual"]
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0]
    assert header[:2] == ("Kode Operasi (wajib)", "Bahan Bakar Aktual (L) (wajib)")
    assert "Alokasi (L) (info)" in header
    codes = {row[0] for row in rows[1:]}
    assert codes == {first["operation_code"], second["operation_code"]}
    assert all(row[1] is None for row in rows[1:])


def test_filled_rows_are_recorded_and_empty_ones_skipped(tmp_path: Path) -> None:
    with TestClient(create_app(database_path=tmp_path / "operations.sqlite3")) as client:
        _train_baseline(client)
        filled = _operation_with_prediction(client, 24)["operation"]
        _operation_with_prediction(client, 36)
        downloaded = load_workbook(BytesIO(client.get("/bahan-bakar-aktual/menunggu.xlsx").content))
        sheet = downloaded["Bahan Bakar Aktual"]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == filled["operation_code"]:
                row[1].value = 21.5
        buffer = BytesIO()
        downloaded.save(buffer)

        response = client.post(
            "/api/v1/bulk-actual-fuel",
            files={
                "file": (
                    "aktual.xlsx",
                    buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["accepted_row_count"] == 1
    assert body["accepted_rows"][0]["actual_fuel"]["operation_id"] == filled["operation_id"]
    assert body["quarantined_row_count"] == 0
    assert body["unfilled_row_count"] == 1


def test_a_code_with_no_litres_in_any_sheet_is_skipped_not_an_error(tmp_path: Path) -> None:
    with TestClient(create_app(database_path=tmp_path / "operations.sqlite3")) as client:
        _train_baseline(client)
        operation = _operation_with_prediction(client, 24)["operation"]
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.append(("Kode Operasi (wajib)", "Bahan Bakar Aktual (L) (wajib)"))
        sheet.append((operation["operation_code"], None))
        buffer = BytesIO()
        workbook.save(buffer)

        body = client.post(
            "/api/v1/bulk-actual-fuel",
            files={"file": ("aktual.xlsx", buffer.getvalue(), "application/octet-stream")},
        ).json()

    assert body["quarantined_row_count"] == 0
    assert body["unfilled_row_count"] == 1


def test_the_sheet_is_for_whoever_records_actual_fuel() -> None:
    assert (
        "GET",
        "/bahan-bakar-aktual/menunggu.xlsx",
        Capability.RECORD_ACTUAL_FUEL,
    ) in ROUTE_CAPABILITIES
