"""Uploaded sheets have a size the server will read, and say so when exceeded.

Every bulk plan, bulk actual-fuel sheet and history import is read whole
into memory, and an .xlsx is a zip that expands as it is read. Nothing
bounded either, while model packages already had limits. Sheets are now
refused above 10 MB or 20.000 rows with a message saying to split them, and
an .xlsx styled down to its last row is not read as a million blank rows.
"""

from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from fuel_predictor.application.historical_datasets import HistoricalDatasetImportError
from fuel_predictor.infrastructure.historical_source_reader import (
    MAX_UPLOAD_BYTES,
    SpreadsheetHistoricalDatasetSourceReader,
)
from fuel_predictor.main import create_app
from tests.test_actual_fuel_outcomes import _train_baseline

_CSV = "Kode Operasi (wajib),Bahan Bakar Aktual (L) (wajib)\n"


def test_an_oversized_file_is_refused_before_it_is_read() -> None:
    reader = SpreadsheetHistoricalDatasetSourceReader(max_bytes=100)

    with pytest.raises(HistoricalDatasetImportError, match="terlalu besar"):
        # Not even a valid sheet: the size is checked before any parsing.
        reader.read("besar.xlsx", b"x" * 101)


def test_a_sheet_over_the_row_limit_asks_to_be_split() -> None:
    reader = SpreadsheetHistoricalDatasetSourceReader(max_rows=3)
    csv = (_CSV + "A,1\n" * 4).encode()

    with pytest.raises(HistoricalDatasetImportError, match="lebih dari 3 baris"):
        reader.read("banyak.csv", csv)
    assert reader.read("cukup.csv", (_CSV + "A,1\n" * 3).encode())[0].rows


def test_the_row_limit_counts_every_sheet_of_a_workbook() -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    second = workbook.create_sheet("Kedua")
    for sheet in (first, second):
        sheet.append(("Kode Operasi", "Bahan Bakar Aktual (L)"))
        for _ in range(2):
            sheet.append(("A", 1))
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(HistoricalDatasetImportError, match="lebih dari 3 baris"):
        SpreadsheetHistoricalDatasetSourceReader(max_rows=3).read("dua.xlsx", buffer.getvalue())


def test_blank_rows_at_the_end_of_an_excel_sheet_are_formatting_not_data() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(("Kode Operasi", "Bahan Bakar Aktual (L)"))
    sheet.append(("A", 1))
    sheet.append((None, None))  # a gap between rows is kept, and reported as blank
    sheet.append(("B", 2))
    # Styled but empty all the way down, as a sheet copied from a template often is.
    sheet.cell(row=50_000, column=1).number_format = "0.00"
    buffer = BytesIO()
    workbook.save(buffer)

    rows = (
        SpreadsheetHistoricalDatasetSourceReader(max_rows=10)
        .read("rapi.xlsx", buffer.getvalue())[0]
        .rows
    )

    assert [number for number, _values in rows] == [2, 3, 4]


def test_the_page_says_the_file_is_too_big(tmp_path: Path) -> None:
    big = _CSV.encode() + b"A,1\n" * (MAX_UPLOAD_BYTES // 4 + 1)
    with TestClient(create_app(database_path=tmp_path / "operations.sqlite3")) as client:
        _train_baseline(client)
        page = client.post(
            "/bahan-bakar-aktual-massal",
            files={"file": ("besar.csv", big, "text/csv")},
        )

    assert page.status_code == 422
    assert "Berkas terlalu besar: maksimal 10 MB" in page.text
