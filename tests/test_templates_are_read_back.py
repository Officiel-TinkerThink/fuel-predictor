"""Every column of a template the app hands out is read back when uploaded.

The bulk prediction template's "Kendaraan (opsional)" normalised to
"kendaraan opsional", which the importer did not know: a sheet filled in from
the app's own template lost every vehicle without a word - no vehicle feature
for the model, no vehicle in the operation code, no lifting check.
"""

from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from fuel_predictor.application.bulk_actual_fuel import _HEADER_ALIASES as ACTUAL_ALIASES
from fuel_predictor.application.bulk_operation_predictions import (
    _HEADER_ALIASES as PREDICTION_ALIASES,
)
from fuel_predictor.application.historical_datasets import normalize_header
from fuel_predictor.infrastructure.actual_fuel_template import BULK_ACTUAL_FUEL_TEMPLATE_HEADERS
from fuel_predictor.infrastructure.bulk_prediction_template import (
    BULK_PREDICTION_TEMPLATE_HEADERS,
)
from fuel_predictor.main import create_app
from tests.test_actual_fuel_outcomes import _operation_with_prediction, _train_baseline


@pytest.mark.parametrize(
    ("headers", "aliases"),
    [
        (BULK_PREDICTION_TEMPLATE_HEADERS, PREDICTION_ALIASES),
        (BULK_ACTUAL_FUEL_TEMPLATE_HEADERS, ACTUAL_ALIASES),
    ],
)
def test_every_template_header_maps_to_a_field(
    headers: tuple[str, ...], aliases: dict[str, set[str]]
) -> None:
    known = {alias for spellings in aliases.values() for alias in spellings}

    assert [header for header in headers if normalize_header(header) not in known] == []


# --- the instructions sheet is not data --------------------------------------------


def _filled(template: bytes, row: tuple[str | float | None, ...]) -> bytes:
    workbook = load_workbook(BytesIO(template))
    workbook[workbook.sheetnames[0]].append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_a_filled_prediction_template_quarantines_nothing(tmp_path: Path) -> None:
    """The template's "Petunjuk" sheet was read as data: one operation filled in
    correctly came back with nine rows "dikarantina" - the instructions."""
    with TestClient(create_app(database_path=tmp_path / "operations.sqlite3")) as client:
        _train_baseline(client)
        template = client.get("/api/v1/bulk-operation-predictions/template?format=xlsx").content
        sheet = _filled(template, (None, "Mobilisasi", 20, None, None))

        page = client.post(
            "/prediksi-operasi-massal",
            files={"file": ("rencana.xlsx", sheet, "application/octet-stream")},
        ).text

    assert "1 operasi mendapat kode dan estimasi." in page
    assert "dikarantina" not in page


def test_a_filled_actual_fuel_template_quarantines_nothing(tmp_path: Path) -> None:
    with TestClient(create_app(database_path=tmp_path / "operations.sqlite3")) as client:
        _train_baseline(client)
        operation = _operation_with_prediction(client, 24)["operation"]
        template = client.get("/api/v1/bulk-actual-fuel/template?format=xlsx").content
        sheet = _filled(template, (operation["operation_code"], 21, "fuel_meter"))

        body = client.post(
            "/api/v1/bulk-actual-fuel",
            files={"file": ("aktual.xlsx", sheet, "application/octet-stream")},
        ).json()

    assert body["accepted_row_count"] == 1
    assert body["quarantined_row_count"] == 0


def test_a_file_with_none_of_the_columns_is_refused_with_a_reason(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(("Nama", "Alamat"))
    sheet.append(("Budi", "Palembang"))
    buffer = BytesIO()
    workbook.save(buffer)
    with TestClient(create_app(database_path=tmp_path / "operations.sqlite3")) as client:
        response = client.post(
            "/api/v1/bulk-actual-fuel",
            files={"file": ("salah.xlsx", buffer.getvalue(), "application/octet-stream")},
        )

    assert response.status_code == 422
    assert "templat" in response.text.lower()
