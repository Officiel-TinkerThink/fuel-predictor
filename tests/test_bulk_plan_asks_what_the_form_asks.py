"""The bulk plan sheet asks what the single form asks, in the form's words.

It used to require "Kategori ANGBER" on every row (every unit is ANGBER),
a "Sumber Jarak" of manual or routing_provider (a choice the form never
offers), activity codes such as transport_and_lifting, and listed six
vehicle names that were long out of date. Now the unit and activity are
dropdowns fed by the fleet, the words are "Mobilisasi" and "Mobilisasi +
lifting", and sheets filled from the old template still import.
"""

from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from fuel_predictor.infrastructure.packaged_vehicle_catalog import PackagedVehicleCatalog
from fuel_predictor.main import create_app
from tests.test_actual_fuel_outcomes import _train_baseline


def _client(tmp_path: Path) -> TestClient:
    return TestClient(
        create_app(
            database_path=tmp_path / "operations.sqlite3", vehicle_catalog=PackagedVehicleCatalog()
        )
    )


def test_the_unit_and_activity_are_dropdowns_fed_by_the_fleet(tmp_path: Path) -> None:
    with _client(tmp_path) as client:
        template = client.get("/api/v1/bulk-operation-predictions/template?format=xlsx").content

    workbook = load_workbook(BytesIO(template))
    names = [row[0] for row in workbook["Daftar unit"].iter_rows(min_row=2, values_only=True)]
    assert "VT 01" in names and "Truck Crane 01" in names and len(names) == 23
    assert workbook["Daftar unit"].sheet_state == "hidden"
    lists = {
        str(rule.sqref): rule.formula1
        for rule in workbook["Operasi Harian"].data_validations.dataValidation
    }
    assert lists["A2:A500"] == "='Daftar unit'!$A$2:$A$24"
    assert lists["B2:B500"] == '"Mobilisasi,Mobilisasi + lifting"'
    guide = " ".join(str(row[2]) for row in workbook["Petunjuk"].iter_rows(values_only=True))
    # The units that may lift, named from the fleet rather than a stale list.
    assert "Truck Crane 01, Truck Crane 02, Wheel Crane" in guide
    assert "Whellcrane" not in guide


def _plan(client: TestClient, *rows: tuple[str | float | None, ...]) -> dict[str, object]:
    template = client.get("/api/v1/bulk-operation-predictions/template?format=xlsx").content
    workbook = load_workbook(BytesIO(template))
    for row in rows:
        workbook["Operasi Harian"].append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    sheet = buffer.getvalue()
    response = client.post(
        "/api/v1/bulk-operation-predictions",
        files={"file": ("rencana.xlsx", sheet, "application/octet-stream")},
    )
    assert response.status_code == 201, response.text
    body: dict[str, object] = response.json()
    return body


def test_the_form_s_words_plan_operations(tmp_path: Path) -> None:
    with _client(tmp_path) as client:
        _train_baseline(client)
        body = _plan(
            client,
            ("VT 01", "Mobilisasi", 42, None, None),
            ("Truck Crane 01", "Mobilisasi + lifting", 28, 3, None),
        )

    assert body["accepted_row_count"] == 2, body
    assert body["quarantined_row_count"] == 0


def test_lifting_on_a_unit_that_cannot_lift_is_refused(tmp_path: Path) -> None:
    with _client(tmp_path) as client:
        _train_baseline(client)
        body = _plan(client, ("VT 01", "Mobilisasi + lifting", 42, 2, None))

    assert body["accepted_row_count"] == 0
    report = body["correction_report"]
    assert isinstance(report, list)
    assert "lifting" in report[0]["reasons"][0]["message"].lower()


def test_an_unknown_activity_says_which_words_to_use(tmp_path: Path) -> None:
    with _client(tmp_path) as client:
        _train_baseline(client)
        body = _plan(client, ("VT 01", "Angkut barang", 42, None, None))

    report = body["correction_report"]
    assert isinstance(report, list)
    assert report[0]["reasons"][0]["message"] == (
        'Aktivitas tidak dikenali; gunakan "Mobilisasi" atau "Mobilisasi + lifting".'
    )


def test_a_sheet_from_the_old_template_still_imports(tmp_path: Path) -> None:
    old = (
        "Kategori ANGBER (wajib),Kendaraan (opsional),Mode Aktivitas (wajib),"
        "Jam Lifting (opsional),Jarak Total (km) (wajib),Sumber Jarak (wajib),"
        "Urutan Pemberhentian (opsional)\n"
        "ANGBER,VT 01,transport,,42,manual,\n"
    )
    with _client(tmp_path) as client:
        _train_baseline(client)
        response = client.post(
            "/api/v1/bulk-operation-predictions",
            files={"file": ("lama.csv", old.encode(), "text/csv")},
        )

    assert response.status_code == 201, response.text
    assert response.json()["accepted_row_count"] == 1


def test_a_unit_written_as_people_write_it_is_planned_under_its_fleet_name(
    tmp_path: Path,
) -> None:
    """Plan rows took the unit as written: "t crane 01" was stored as a unit
    of its own, which the model had never seen, while the history import
    already read it as Truck Crane 01."""
    with _client(tmp_path) as client:
        _train_baseline(client)
        body = _plan(
            client,
            ("t crane 01", "Mobilisasi + lifting", 28, 3, None),
            ("PM 01", "Mobilisasi", 30, None, None),
        )

    accepted = body["accepted_rows"]
    assert isinstance(accepted, list)
    assert [row["operation"]["vehicle"] for row in accepted] == ["Truck Crane 01", "Prime Mover"]


def test_a_unit_the_fleet_does_not_know_is_planned_as_written(tmp_path: Path) -> None:
    """A rented crane is planned before it is ever catalogued, as on the API."""
    with _client(tmp_path) as client:
        _train_baseline(client)
        body = _plan(client, ("Crane Sewa 01", "Mobilisasi", 30, None, None))

    accepted = body["accepted_rows"]
    assert isinstance(accepted, list)
    assert [row["operation"]["vehicle"] for row in accepted] == ["Crane Sewa 01"]
